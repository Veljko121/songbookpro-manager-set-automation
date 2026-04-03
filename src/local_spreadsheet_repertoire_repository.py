from base_repertoire_repository import BaseRepertoireRepository
import openpyxl as xl

class LocalSpreadsheetRepertoireRepository(BaseRepertoireRepository):

    def __init__(self):
        super().__init__()

    def get_sheets(self, spreadsheet_id: str):
        return xl.load_workbook(spreadsheet_id, read_only=True).sheetnames
    
    def _fetch_song_data(self, spreadsheet_id, sheet_id, song_names_column, keys_column, notes_column):
        doc = xl.load_workbook(spreadsheet_id)
        sheet = doc[sheet_id]

        def get_column_values(col_index):
            return [cell.value for cell in sheet[xl.utils.get_column_letter(col_index)]]

        song_names = get_column_values(song_names_column)
        song_keys = get_column_values(keys_column)
        song_notes = get_column_values(notes_column)

        return song_names, song_keys, song_notes
    
if __name__ == "__main__":
    spreadsheet_path = "/home/veljko/Downloads/piknik.xlsx"
    repo = LocalSpreadsheetRepertoireRepository()
    print(repo.get_songs(spreadsheet_path, "Piknik", 1, 2, 3))